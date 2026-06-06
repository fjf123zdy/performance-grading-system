#!/usr/bin/env python3
"""
生成合成革企业环境绩效分级评价系统 V1.0 的软件著作权申请文档：
  1. 用户手册 (.docx)
  2. 程序代码文档 (.docx)
"""

import os
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

# ============================================================
# 路径配置
# ============================================================
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
SRC_DIR = os.path.join(BASE_DIR, 'src')
OUTPUT_DIR = BASE_DIR

# ============================================================
# 通用样式函数
# ============================================================
def set_cell_font(cell, name='宋体', size=10.5, bold=False, color=None):
    """设置单元格字体"""
    for paragraph in cell.paragraphs:
        for run in paragraph.runs:
            run.font.name = name
            run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
            run.font.size = Pt(size)
            run.bold = bold
            if color:
                run.font.color.rgb = RGBColor(*color)

def add_heading_styled(doc, text, level=1):
    """添加带样式的标题"""
    heading = doc.add_heading(text, level=level)
    for run in heading.runs:
        run.font.name = '黑体'
        run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    return heading

def add_para(doc, text, bold=False, font_name='宋体', size=12, alignment=None, indent=True):
    """添加段落"""
    para = doc.add_paragraph()
    if indent:
        para.paragraph_format.first_line_indent = Pt(size * 2)
    if alignment is not None:
        para.alignment = alignment
    run = para.add_run(text)
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    run.font.size = Pt(size)
    run.bold = bold
    return para

def add_code_block(doc, code_text, font_size=8):
    """添加代码块（等宽字体）"""
    para = doc.add_paragraph()
    para.paragraph_format.first_line_indent = Pt(0)
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.line_spacing = 1.0
    run = para.add_run(code_text)
    run.font.name = 'Courier New'
    run.font.size = Pt(font_size)
    return para

def add_table_with_data(doc, headers, rows, col_widths=None):
    """添加表格"""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # 表头
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_font(cell, name='黑体', size=10, bold=True)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # 数据行
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = table.rows[r + 1].cells[c]
            cell.text = str(val)
            set_cell_font(cell, name='宋体', size=10)

    return table

# ============================================================
# 文档一：用户手册
# ============================================================
def create_user_manual():
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(12)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # ---- 封面 ----
    doc.add_paragraph()
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('合成革企业环境绩效分级评价系统 V1.0')
    run.font.name = '黑体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    run.font.size = Pt(22)
    run.bold = True

    doc.add_paragraph()
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('用户手册')
    run.font.name = '黑体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    run.font.size = Pt(18)

    doc.add_paragraph()
    doc.add_paragraph()

    cover_info = [
        '软件版本：V1.0',
        '开发日期：2026年6月',
        '文档类型：软件著作权申请 — 用户手册',
        '开发语言：JavaScript (Vue 3)',
        '运行环境：浏览器 (Chrome/Firefox/Edge/Safari)'
    ]
    for info in cover_info:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(info)
        run.font.size = Pt(14)

    doc.add_page_break()

    # ---- 目录 ----
    add_heading_styled(doc, '目  录', level=1)
    toc_items = [
        '一、引言',
        '  1.1 编制目的',
        '  1.2 适用范围',
        '  1.3 参考资料',
        '二、系统概述',
        '  2.1 系统功能架构',
        '  2.2 运行环境要求',
        '  2.3 系统启动方式',
        '三、功能模块操作说明',
        '  3.1 系统首页',
        '  3.2 企业管理模块',
        '  3.3 基本要求判定模块',
        '  3.4 定量指标计算模块',
        '  3.5 定性指标评分模块',
        '  3.6 AHP 权重计算模块',
        '  3.7 PSO 优化模块',
        '  3.8 群决策融合模块',
        '  3.9 分级结果展示模块',
        '  3.10 综合报告模块',
        '四、评价流程说明',
        '五、分级结果应用',
        '六、注意事项'
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Pt(0)
        run = p.add_run(item)
        run.font.size = Pt(12)
        if not item.startswith('  '):
            run.bold = True

    doc.add_page_break()

    # ---- 一、引言 ----
    add_heading_styled(doc, '一、引言', level=1)

    add_heading_styled(doc, '1.1 编制目的', level=2)
    add_para(doc, '本用户手册旨在为"合成革企业环境绩效分级评价系统 V1.0"的使用者提供完整的操作指导。该系统基于中国环境科学研究院编制的《合成革企业环境管理与排放水平绩效分级政策建议研究报告》（2025年12月），将合成革企业环保绩效分级指标体系数字化，实现对合成革企业环境管理与排放水平的科学量化评价和自动分级。')
    add_para(doc, '通过本手册，用户可了解系统的功能架构、操作流程和注意事项，快速掌握系统使用方法，完成对合成革企业的环境绩效分级评价工作。')

    add_heading_styled(doc, '1.2 适用范围', level=2)
    add_para(doc, '本系统适用于合成革行业企业的环境管理与排放水平绩效分级评价。适用对象包括：各级生态环境管理部门、合成革生产企业、环境评价机构、行业研究单位等。')
    add_para(doc, '系统评价结果可作为重污染天气应急减排措施制定、企业产能调控、环保税收优惠等政策执行的参考依据。')

    add_heading_styled(doc, '1.3 参考资料', level=2)
    add_para(doc, '（1）《合成革企业环境管理与排放水平绩效分级政策建议研究报告》，中国环境科学研究院，2025年12月。')
    add_para(doc, '（2）《国家重污染天气重点行业应急减排措施制定技术指南》。')
    add_para(doc, '（3）《中共中央国务院关于全面推进美丽中国建设的意见》。')
    add_para(doc, '（4）Saaty, T.L. (1980) The Analytic Hierarchy Process.')
    add_para(doc, '（5）Kennedy, J. & Eberhart, R. (1995) Particle Swarm Optimization.')

    # ---- 二、系统概述 ----
    add_heading_styled(doc, '二、系统概述', level=1)

    add_heading_styled(doc, '2.1 系统功能架构', level=2)
    add_para(doc, '系统采用 B/S（浏览器/服务器）架构设计，前端基于 Vue 3 框架开发，通过浏览器即可访问使用。系统功能架构分为以下几个层次：')
    add_para(doc, '（1）数据层：管理企业基本信息和模拟数据，维护5家预设企业和5位专家打分数据。')
    add_para(doc, '（2）业务层：实现基本要求判定、定量指标计算、定性指标评分、AHP权重分析、PSO算法优化、群决策融合、综合分级判定等核心业务逻辑。')
    add_para(doc, '（3）展示层：提供系统首页、企业管理、各评价模块页面、分级结果展示和综合报告等10个页面视图。')

    add_heading_styled(doc, '2.2 运行环境要求', level=2)
    add_para(doc, '本系统为纯前端 Web 应用，运行环境要求如下：')
    add_para(doc, '硬件环境：CPU 1GHz以上，内存 2GB以上，显示器分辨率 1366×768及以上。')
    add_para(doc, '软件环境：操作系统 Windows/macOS/Linux，浏览器 Chrome 90+ / Firefox 88+ / Edge 90+ / Safari 14+。')
    add_para(doc, '开发环境：Node.js 18+，npm 9+，Vite 5+。')

    add_heading_styled(doc, '2.3 系统启动方式', level=2)
    add_para(doc, '本系统提供两种运行方式：')
    add_para(doc, '方式一（开发模式）：在项目根目录执行"npm install"安装依赖，再执行"npm run dev"启动开发服务器，浏览器访问 http://localhost:5173。')
    add_para(doc, '方式二（生产模式）：执行"npm run build"构建生产版本，将 dist/ 目录部署至任意静态文件服务器，通过服务器地址访问。')

    # ---- 三、功能模块操作说明 ----
    add_heading_styled(doc, '三、功能模块操作说明', level=1)

    # 3.1
    add_heading_styled(doc, '3.1 系统首页', level=2)
    add_para(doc, '系统首页是用户进入系统后的第一个页面，顶部导航栏显示系统名称和当前时间，左侧为功能导航菜单，右侧为内容展示区。')
    add_para(doc, '首页中间区域展示系统的6大核心功能模块入口（企业管理、定量指标、定性指标、AHP权重、PSO优化、分级结果），每个模块入口以卡片形式呈现，用户可点击相应卡片快速进入功能页面。')
    add_para(doc, '首页底部为"快速开始"引导区域，列出从选择企业到查看分级结果的4个步骤，帮助新用户快速上手使用系统。页面顶部设有"开始评价"按钮，点击后可直接进入企业管理页面。')

    # 3.2
    add_heading_styled(doc, '3.2 企业管理模块', level=2)
    add_para(doc, '企业管理模块用于管理合成革企业的基本信息，是评价流程的起点。该模块以卡片网格形式展示所有企业信息，每张卡片包含企业名称、所在地、主要产品、年产能、员工人数等基本信息，以及评价状态标签（待评价/已评价）。')
    add_para(doc, '主要操作功能：')
    add_para(doc, '（1）选择企业：点击企业卡片即可选中当前操作企业，选中后卡片会高亮显示。')
    add_para(doc, '（2）开始评价：每张企业卡片底部提供"基本判定""定量指标""定性指标"三个快捷入口按钮，点击即可进入对应评价环节。已完成评价的企业还会显示"查看结果"按钮。')
    add_para(doc, '（3）新增企业：点击页面右上角"+ 新增企业"按钮，弹出新增企业表单弹窗。填写企业名称、省份、城市、主要产品、年产能、员工人数等必要信息后，点击"确认新增"即可完成企业创建。系统将自动为新企业生成评价所需的全部数据（基本要求审查数据、产品折算数据、排放数据、定性指标评分数据），新企业可直接参与全流程评价。')
    add_para(doc, '（4）删除企业：每张企业卡片右上角设有"删除"按钮，点击后弹出确认对话框，确认后系统将同时删除该企业及其关联的所有评价记录。')

    # 3.3
    add_heading_styled(doc, '3.3 基本要求判定模块', level=2)
    add_para(doc, '基本要求判定模块对企业是否具备参与绩效分级的基本资格进行"一票否决制"审查。审查项目包括五项：近三年无环境事故、近三年无安全事故、近三年无质量事故、未列入失信名录、废水排口设置合规。')
    add_para(doc, '用户在企业卡片中点击"基本判定"按钮进入该模块。系统自动读取企业基本审查数据，以表格形式展示每项审查的判定结果和佐证材料。全部通过时显示绿色"通过"标识，可继续后续评价流程；任一项不通过则显示红色"不通过"标识及相关原因说明，该企业不具备参评资格。')

    # 3.4
    add_heading_styled(doc, '3.4 定量指标计算模块', level=2)
    add_para(doc, '定量指标计算模块实现对企业环保绩效定量指标的科学计算和分级判定。该模块包含两个主要部分：')
    add_para(doc, '（1）标准品产量折算：以厚度1.0mm、幅宽136-156cm为标准品基准，根据企业各产品的实际厚度和幅宽查询折算系数，计算标准品总产量（Qbz）。折算结果以表格形式展示各产品的折算系数和标准品产量。')
    add_para(doc, '（2）定量指标计算与分级：基于标准品产量，计算四项定量指标：单位产品废水产生量（Vci=Vc/Qbz）、单位产品VOCs产生量（Gvoc/Qbz）、有组织排放DMF浓度、非封闭区域VOC检测值。每项指标根据研究报告设定的三级阈值进行分级判定（一级/二级/三级/不达标）。')
    add_para(doc, '计算结果表格中包含计算值、各级阈值对照和判定结果，底部显示综合定量等级。')

    # 3.5
    add_heading_styled(doc, '3.5 定性指标评分模块', level=2)
    add_para(doc, '定性指标评分模块对企业在6个管理维度的表现进行综合评分。6个一级指标分别为：原辅材料准备、废气治理、废水收集和处理、初期雨水收集和处理、固废管理、检测监控要求，共包含15个二级指标。')
    add_para(doc, '每项二级指标按照0-10分制进行打分，系统自动读取企业数据并显示各指标得分。得分等级分为四级：优秀（8-10分）、良好（6-7分）、一般（4-5分）、较差（0-3分）。')
    add_para(doc, '页面底部展示评分汇总区域，包含总分（满分10分制）、百分制得分和对应的等级判定结果。')

    # 3.6
    add_heading_styled(doc, '3.6 AHP 权重计算模块', level=2)
    add_para(doc, 'AHP（层次分析法）权重计算模块用于确定评价指标体系中各指标的权重。用户可在此模块选择不同的专家，查看其基于Saaty 1-9标度法构造的判断矩阵及AHP分析结果。')
    add_para(doc, '模块展示内容包括：6×6一级指标判断矩阵（可编辑表格形式）、和积法计算的权重向量（含百分比柱状图）、一致性检验结果（λmax、CI、RI、CR值）。')
    add_para(doc, '当CR<0.1时，判断矩阵通过一致性检验，显示绿色"通过"标识；当CR≥0.1时，显示红色"不通过"标识，提示用户需要进行PSO优化修正。')

    # 3.7
    add_heading_styled(doc, '3.7 PSO 优化模块', level=2)
    add_para(doc, 'PSO（粒子群优化）模块用于对不一致的判断矩阵（CR≥0.1）进行自动修正。用户可在该模块调整PSO算法的运行参数（种群规模、最大迭代次数、α-CR权重），点击"运行PSO优化"按钮执行优化计算。')
    add_para(doc, '优化结果页面展示原始CR值与优化后CR值的对比、优化后的权重向量、迭代次数和最终适应度值。通过PSO算法在最小修改原始矩阵的前提下，使修正后的判断矩阵满足一致性要求（CR<0.1）。')

    # 3.8
    add_heading_styled(doc, '3.8 群决策融合模块', level=2)
    add_para(doc, '群决策融合模块用于综合多位专家的权重分析结果，形成最终的评价指标权重体系。用户可以自主选择参与融合的专家（至少2位），并选择融合方法。')
    add_para(doc, '系统提供两种融合方法：')
    add_para(doc, '（1）直接均值法：对各位专家的权重向量取算术平均，简单直观。')
    add_para(doc, '（2）群决策矩阵法：先对多位专家的判断矩阵各元素取几何平均，构造群决策矩阵，再执行AHP分析。如果群决策矩阵不满足一致性，系统可自动调用PSO算法进行优化修正。')
    add_para(doc, '融合结果以表格形式展示各一级指标的最终权重及其百分比柱状图。')

    # 3.9
    add_heading_styled(doc, '3.9 分级结果展示模块', level=2)
    add_para(doc, '分级结果展示模块是评价流程的终点，综合基本要求判定、定量指标计算和定性指标评分的结果，给出企业最终的环保绩效等级。')
    add_para(doc, '页面核心展示内容包括：')
    add_para(doc, '（1）等级徽章：以金色（引领级）、银色（先进级）、铜色（基础级）圆形徽章动画展示企业最终等级，徽章下方显示等级名称和描述。')
    add_para(doc, '（2）基本要求审查摘要：展示五项审查的通过/不通过状态。')
    add_para(doc, '（3）定量指标详情表：列出四项指标的实测值、各级阈值、单项判定等级。')
    add_para(doc, '（4）管控措施建议：根据企业最终等级，展示差异化的管控措施建议，包括重污染天气应急措施、产能与扩建规定、提标改造要求、监管频次等内容。')

    # 3.10
    add_heading_styled(doc, '3.10 综合报告模块', level=2)
    add_para(doc, '综合报告模块以格式化文档的形式汇总企业环境绩效分级评价的完整结果。报告内容包括：')
    add_para(doc, '（1）企业基本信息（名称、所在地、主要产品、产能、员工人数等）。')
    add_para(doc, '（2）评价结果概览（最终等级、定量等级、定性评分百分比）。')
    add_para(doc, '（3）定量指标详情表（含各级阈值对照）。')
    add_para(doc, '（4）管控措施建议。')
    add_para(doc, '（5）报告自动生成时间和评价依据说明。')
    add_para(doc, '报告可在浏览器中直接查看，适合作为正式评价报告的电子版本存档。')

    # ---- 四、评价流程说明 ----
    add_heading_styled(doc, '四、评价流程说明', level=1)
    add_para(doc, '系统对合成革企业的环境绩效分级评价遵循以下标准流程：')
    add_para(doc, '第一步：选择企业。在企业管理页面中选中待评价企业，或新增一家企业。')
    add_para(doc, '第二步：基本要求判定。系统对企业五项基本资格进行一票否决制审查。不通过则评价终止。')
    add_para(doc, '第三步：定量指标计算。系统对标准品产量进行折算，计算四项定量指标并逐项分级。')
    add_para(doc, '第四步：定性指标评分。对15项二级指标进行0-10分制评分，计算加权总分。')
    add_para(doc, '第五步：权重分析。在AHP模块中计算指标权重，必要时使用PSO修正，通过群决策融合确定最终权重体系。')
    add_para(doc, '第六步：查看分级结果。在分级结果页面查看企业最终等级及对应的管控措施建议。')
    add_para(doc, '第七步：生成综合报告。在综合报告页面查看并导出完整评价报告。')
    add_para(doc, '整个评价流程的进度在页面顶部的步骤指示器中实时显示，用户可点击已完成步骤快速跳转。')

    # ---- 五、分级结果应用 ----
    add_heading_styled(doc, '五、分级结果应用', level=1)
    add_para(doc, '根据《合成革企业环境管理与排放水平绩效分级政策建议研究报告》第四章的规定，不同等级企业适用差异化的管控措施：')

    add_para(doc, '一级企业（引领级）：黄色及以上预警期间，企业可自主采取减排措施；持续半小时以上的小到中雨天气，企业可自主采取减排措施；地方政府可据此审批扩大产能，并优先给予地方税收、补贴等相关政策。')

    add_para(doc, '二级企业（先进级）：黄色及橙色预警期间，企业停产50%；红色预警期间，企业全面停产；持续半小时以上的小到中雨天气，政企协商，建议企业减产50%；企业保持现有产能，不得扩大产能。')

    add_para(doc, '三级企业（基础级）：黄色及橙色预警期间，企业停产50%；红色预警期间，企业全面停产；持续半小时以上的小到中雨天气，政企协商，建议企业停产至雨水天气停止；连续两年获得三级的企业，应提出整改方案并按期整改，逾期不改将调整缩减产能。')

    # ---- 六、注意事项 ----
    add_heading_styled(doc, '六、注意事项', level=1)
    add_para(doc, '（1）本系统为演示原型版本，使用模拟数据进行测试，未连接真实数据库。所有企业数据和专家打分数据均为模拟数据。')
    add_para(doc, '（2）系统中预设的5家企业和5位专家数据仅供系统功能演示和开发测试使用。')
    add_para(doc, '（3）定性指标评分中使用的权重默认为均权，进行完整的AHP+PSO+群决策流程后可使用融合权重重新计算评分。')
    add_para(doc, '（4）系统构建后生成的dist/目录可部署至任意静态文件服务器。')
    add_para(doc, '（5）本系统仅依赖4个运行时库（Vue 3、Vue Router、Pinia、ECharts），所有算法（AHP、PSO、群决策）均为原始实现。')

    # 保存
    output_path = os.path.join(OUTPUT_DIR, '合成革企业环境绩效分级评价系统V1.0_用户手册.docx')
    doc.save(output_path)
    print(f'用户手册已保存至：{output_path}')
    return output_path


# ============================================================
# 文档二：程序代码文档
# ============================================================
def create_source_code_doc():
    doc = Document()

    # 设置默认字体
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(10.5)
    style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

    # ---- 封面 ----
    doc.add_paragraph()
    doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('合成革企业环境绩效分级评价系统 V1.0')
    run.font.name = '黑体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    run.font.size = Pt(22)
    run.bold = True

    doc.add_paragraph()
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run('程序代码文档')
    run.font.name = '黑体'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
    run.font.size = Pt(18)

    doc.add_paragraph()
    doc.add_paragraph()

    code_cover_info = [
        '软件版本：V1.0',
        '开发日期：2026年6月',
        '文档类型：软件著作权申请 — 程序代码文档',
        '开发语言：JavaScript（Vue 3 框架）',
        '源代码文件数：36个',
        '源代码总行数：5,522行'
    ]
    for info in code_cover_info:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(info)
        run.font.size = Pt(14)

    doc.add_page_break()

    # ---- 一、源代码总体说明 ----
    add_heading_styled(doc, '一、源代码总体说明', level=1)
    add_para(doc, '本程序代码文档对应"合成革企业环境绩效分级评价系统 V1.0"的全部原创源代码。系统采用 Vue 3 + Vite 前端框架开发，使用 JavaScript 作为主要开发语言。', size=10.5)
    add_para(doc, f'源代码由36个文件组成，总计5,522行有效代码。文件分布为：JavaScript服务层（8个文件）、Vue视图层（10个页面+3个布局组件）、JavaScript工具层（4个文件）、状态管理层（3个文件）、路由配置（1个文件）、CSS样式层（5个文件）、应用入口（2个文件）。', size=10.5)
    add_para(doc, '所有核心算法（AHP层次分析法、PSO粒子群优化算法、群决策融合算法）均为基于原始学术论文公式的原创实现，未引用任何现成算法库。', size=10.5)

    # ---- 二、源代码文件清单 ----
    add_heading_styled(doc, '二、源代码文件清单', level=1)

    # 收集所有源文件
    src_files = []
    for root, dirs, files in os.walk(SRC_DIR):
        for f in files:
            if f.endswith(('.js', '.vue', '.css')):
                full_path = os.path.join(root, f)
                rel_path = os.path.relpath(full_path, BASE_DIR)
                lines = sum(1 for _ in open(full_path, 'r', encoding='utf-8'))
                src_files.append((rel_path, lines, f))

    src_files.sort(key=lambda x: x[0])

    # 文件清单表
    add_table_with_data(
        doc,
        ['序号', '文件路径', '文件名', '代码行数', '文件类型'],
        [[i+1, f[0], f[2], str(f[1]), f[2].split('.')[-1].upper()] for i, f in enumerate(src_files)]
    )

    doc.add_paragraph()
    add_para(doc, f'合计：{len(src_files)} 个文件，{sum(f[1] for f in src_files)} 行源代码。', size=10.5)

    doc.add_page_break()

    # ---- 三、源代码正文（按模块排列） ----
    add_heading_styled(doc, '三、源代码正文', level=1)
    add_para(doc, '以下按照系统模块结构顺序，列出全部源代码文件的内容。代码中包含完整的中文注释，标注了算法依据、计算公式和数据结构的说明。', size=10.5)

    # 按模块分组读取源代码
    modules = [
        ('应用入口', ['src/main.js', 'src/App.vue', 'src/router/index.js']),
        ('CSS 样式', ['src/assets/styles/variables.css', 'src/assets/styles/main.css', 'src/assets/styles/form.css', 'src/assets/styles/table.css', 'src/assets/styles/badge.css']),
        ('工具函数', ['src/utils/constants.js', 'src/utils/mathHelper.js', 'src/utils/validators.js', 'src/utils/formatters.js']),
        ('状态管理', ['src/stores/useEnterpriseStore.js', 'src/stores/useEvaluationStore.js', 'src/stores/useExpertStore.js']),
        ('核心服务 - 模拟数据', ['src/services/mockData.js']),
        ('核心服务 - 业务逻辑', ['src/services/basicCheckService.js', 'src/services/quantitativeService.js', 'src/services/qualitativeService.js', 'src/services/gradingService.js']),
        ('核心服务 - 算法引擎', ['src/services/ahpService.js', 'src/services/psoService.js', 'src/services/groupDecisionService.js']),
        ('布局组件', ['src/components/layout/AppHeader.vue', 'src/components/layout/AppSidebar.vue', 'src/components/layout/StepIndicator.vue']),
        ('页面视图', ['src/views/HomeView.vue', 'src/views/EnterpriseView.vue', 'src/views/BasicCheckView.vue', 'src/views/QuantitativeView.vue', 'src/views/QualitativeView.vue', 'src/views/AHPView.vue', 'src/views/PSOView.vue', 'src/views/GroupDecisionView.vue', 'src/views/GradingView.vue', 'src/views/ReportView.vue']),
    ]

    for module_name, file_list in modules:
        add_heading_styled(doc, f'模块：{module_name}', level=2)

        for rel_path in file_list:
            full_path = os.path.join(BASE_DIR, rel_path)
            if not os.path.exists(full_path):
                continue

            # 文件名标题
            add_heading_styled(doc, f'文件：{rel_path}', level=3)

            # 读取并输出代码
            try:
                with open(full_path, 'r', encoding='utf-8') as f:
                    code_lines = f.readlines()

                # 文件信息
                add_para(doc, f'（共 {len(code_lines)} 行）', size=9, indent=False)

                # 输出代码内容（每行一个代码块段落）
                for line in code_lines:
                    # 去掉行尾换行符，保留缩进
                    line_text = line.rstrip('\n\r')
                    # 将制表符转为空格以便显示
                    display_line = line_text.replace('\t', '    ')
                    # 限制单行长度
                    if len(display_line) > 120:
                        display_line = display_line[:117] + '...'
                    add_code_block(doc, display_line, font_size=7)

            except Exception as e:
                add_para(doc, f'[读取文件出错：{str(e)}]', size=9, indent=False)

    # 保存
    output_path = os.path.join(OUTPUT_DIR, '合成革企业环境绩效分级评价系统V1.0_程序代码文档.docx')
    doc.save(output_path)
    print(f'程序代码文档已保存至：{output_path}')
    return output_path


# ============================================================
# 主程序
# ============================================================
if __name__ == '__main__':
    print('开始生成软件著作权申请文档...\n')

    path1 = create_user_manual()
    print()
    path2 = create_source_code_doc()

    print(f'\n===== 文档生成完成 =====')
    print(f'用户手册：{path1}')
    print(f'程序代码文档：{path2}')
